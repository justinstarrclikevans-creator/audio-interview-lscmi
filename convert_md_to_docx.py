#!/usr/bin/env python3
"""
Convert Markdown documents to beautifully styled Microsoft Word (.docx) documents.
Supports headings, bold/italics, bullet/numbered lists, blockquotes, tables,
and genuine clickable hyperlinks for jobs, referrals, and web resources.
"""

import sys
import os
import re
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_hyperlink(paragraph, url, text, color="0284C7", underline=True):
    """Add a clickable Word hyperlink to a paragraph."""
    part = paragraph.part
    r_id = part.relate_to(url, docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK, is_external=True)

    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), r_id)

    new_run = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')

    if color:
        c = OxmlElement('w:color')
        c.set(qn('w:val'), color)
        rPr.append(c)

    if underline:
        u = OxmlElement('w:u')
        u.set(qn('w:val'), 'single')
        rPr.append(u)

    new_run.append(rPr)
    new_run.text = text
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)

def render_formatted_spans(p, text, is_blockquote=False):
    """Parse text and render bold, italics, code, and clickable hyperlinks."""
    pattern = re.compile(r'(\[[^\]]+\]\([^\)]+\)|\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)')
    last_end = 0

    for m in pattern.finditer(text):
        start, end = m.span()
        if start > last_end:
            plain_chunk = text[last_end:start]
            run = p.add_run(plain_chunk)
            if is_blockquote:
                run.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

        token = m.group(0)
        if token.startswith('[') and token.endswith(')'):
            m_link = re.match(r'^\[(.*?)\]\((.*?)\)$', token)
            if m_link:
                link_text = m_link.group(1)
                link_url = m_link.group(2)
                add_hyperlink(p, link_url, link_text, color="0284C7", underline=True)
        elif token.startswith('**') and token.endswith('**'):
            run = p.add_run(token[2:-2])
            run.bold = True
            if is_blockquote:
                run.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
        elif token.startswith('*') and token.endswith('*'):
            run = p.add_run(token[1:-1])
            run.italic = True
        elif token.startswith('`') and token.endswith('`'):
            run = p.add_run(token[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(0x0f, 0x17, 0x2a)
        last_end = end

    if last_end < len(text):
        plain_chunk = text[last_end:]
        run = p.add_run(plain_chunk)
        if is_blockquote:
            run.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

def add_styled_paragraph(doc, text, style='Normal', is_blockquote=False):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15

    if is_blockquote:
        p.paragraph_format.left_indent = Inches(0.4)
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)

    render_formatted_spans(p, text, is_blockquote=is_blockquote)
    return p

def convert_md_to_docx(md_path, docx_path):
    if not os.path.exists(md_path):
        print(f"File not found: {md_path}")
        return False

    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    doc = docx.Document()

    # Set page margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)

    # Set base Normal style font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)
    font.color.rgb = RGBColor(0x1e, 0x29, 0x3b)

    in_table = False
    table_lines = []

    def flush_table():
        nonlocal table_lines, in_table
        if not table_lines:
            return
        
        parsed_rows = []
        for line in table_lines:
            line_str = line.strip()
            if re.match(r'^\|?[\s\-:|]+\|?$', line_str):
                continue
            cols = [c.strip() for c in line_str.strip('|').split('|')]
            if cols:
                parsed_rows.append(cols)

        if parsed_rows:
            num_cols = max(len(r) for r in parsed_rows)
            table = doc.add_table(rows=len(parsed_rows), cols=num_cols)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = True

            for r_idx, row_data in enumerate(parsed_rows):
                row = table.rows[r_idx]
                for c_idx in range(num_cols):
                    cell = row.cells[c_idx]
                    val = row_data[c_idx] if c_idx < len(row_data) else ""
                    cell.text = ""
                    cp = cell.paragraphs[0]
                    cp.paragraph_format.space_after = Pt(2)
                    cp.paragraph_format.space_before = Pt(2)
                    
                    if r_idx == 0:
                        set_cell_background(cell, "1E3A8A")
                        run = cp.add_run(val)
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                    else:
                        bg = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
                        set_cell_background(cell, bg)
                        render_formatted_spans(cp, val)

                    set_cell_margins(cell, top=80, bottom=80, left=120, right=120)

            sp = doc.add_paragraph()
            sp.paragraph_format.space_before = Pt(4)
            sp.paragraph_format.space_after = Pt(4)

        table_lines = []
        in_table = False

    for line in lines:
        stripped = line.strip()

        # Check for table line
        if stripped.startswith('|') or (in_table and '|' in stripped):
            in_table = True
            table_lines.append(stripped)
            continue
        elif in_table:
            flush_table()

        if not stripped:
            continue

        # Headings
        if stripped.startswith('# '):
            h = doc.add_heading(level=1)
            h.paragraph_format.space_before = Pt(14)
            h.paragraph_format.space_after = Pt(6)
            run = h.add_run(stripped[2:])
            run.font.name = 'Calibri'
            run.font.size = Pt(18)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
        elif stripped.startswith('## '):
            h = doc.add_heading(level=2)
            h.paragraph_format.space_before = Pt(12)
            h.paragraph_format.space_after = Pt(4)
            run = h.add_run(stripped[3:])
            run.font.name = 'Calibri'
            run.font.size = Pt(14)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
        elif stripped.startswith('### '):
            h = doc.add_heading(level=3)
            h.paragraph_format.space_before = Pt(10)
            h.paragraph_format.space_after = Pt(3)
            run = h.add_run(stripped[4:])
            run.font.name = 'Calibri'
            run.font.size = Pt(12)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x02, 0x84, 0xC7)
        elif stripped.startswith('#### '):
            h = doc.add_heading(level=4)
            h.paragraph_format.space_before = Pt(8)
            h.paragraph_format.space_after = Pt(2)
            run = h.add_run(stripped[5:])
            run.font.name = 'Calibri'
            run.font.size = Pt(11)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
        # Bullet list
        elif re.match(r'^[-*+]\s+', stripped):
            item_text = re.sub(r'^[-*+]\s+', '', stripped)
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            render_formatted_spans(p, item_text)
        # Numbered list
        elif re.match(r'^\d+\.\s+', stripped):
            item_text = re.sub(r'^\d+\.\s+', '', stripped)
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            render_formatted_spans(p, item_text)
        # Blockquote
        elif stripped.startswith('>'):
            quote_text = stripped.lstrip('> ').strip()
            add_styled_paragraph(doc, quote_text, is_blockquote=True)
        # Regular paragraph
        else:
            add_styled_paragraph(doc, stripped)

    if in_table:
        flush_table()

    doc.save(docx_path)
    print(f"✅ Generated DOCX with clickable hyperlinks: {os.path.basename(docx_path)}")
    return True

if __name__ == '__main__':
    if len(sys.argv) < 3:
        base_dir = os.path.dirname(os.path.abspath(__file__))
        data_dir = os.path.join(base_dir, 'audio-interview-app', 'data')
        outputs_dir = os.path.join(base_dir, 'reentry-nav-desktop', 'outputs')
        
        count = 0
        for target_dir in [data_dir, outputs_dir]:
            if os.path.exists(target_dir):
                for f in os.listdir(target_dir):
                    if f.endswith('.md'):
                        md = os.path.join(target_dir, f)
                        docx_file = os.path.join(target_dir, f.replace('.md', '.docx'))
                        convert_md_to_docx(md, docx_file)
                        count += 1
        print(f"\nBatch conversion complete: {count} DOCX files created.")
    else:
        convert_md_to_docx(sys.argv[1], sys.argv[2])
